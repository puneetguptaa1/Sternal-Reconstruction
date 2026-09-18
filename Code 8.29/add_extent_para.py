#!/usr/bin/env python3
"""Add a short resection-extent paragraph to the oncological Results section,
with real RefWorks citation controls.

Every reference used here is already cited elsewhere in the manuscript, so no
new item enters the RefWorks store and the bibliography does not renumber. The
displayed numbers were read from the document's own bibliography and
independently cross-checked against the number sets of the existing citation
controls, so they are correct as written.
"""
import copy, json, random, sys, docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
sys.path.insert(0, "/tmp/doc")
from docxlib import annotate, set_text, citation_nodes

SRC, DST = "ms27.docx", "ms28.docx"
REFS = json.load(open("onco_refs.json"))

def pick(*keys):
    out = []
    for k in keys:
        m = [v for kk, v in REFS.items() if kk.startswith(k)]
        assert len(m) == 1, (k, m)
        out.append(m[0])
    return out

RIBS = pick("Asanuma", "Bergovec", "Danino", "Dingemann", "Galbis",
            "Grosfeld", "Jonsson", "Ma et al", "Nishida", "Novoa",
            "Scarnecchia", "Spicer")
AREA = pick("Asanuma", "Berthet", "D'Amico", "Hanna", "Haraguchi",
            "Huang", "Ma et al", "Tasnim")
SPICER = pick("Spicer")

def render(nums):
    """RefWorks collapses runs of three or more consecutive numbers."""
    nums = sorted(nums)
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(f"{nums[i]}-{nums[j]}" if j - i >= 2 else
                     ", ".join(str(n) for n in nums[i:j + 1]))
        i = j + 1
    return ", ".join(parts)

OPT = {"author": True, "year": True, "formatAuthorYear": False,
       "pageReplace": "", "additionalField": "", "additionalValue": "",
       "prefix": "", "suffix": ""}
SPAN = ('<span style="font-family:Times New Roman;font-size:'
        '13.333333333333332px;color:#000000"><sup>%s</sup></span>')

def make_citation(template, group):
    """A new RefWorks citation control carrying `group`'s references."""
    ids = [g["doc"] for g in group]
    text = render([g["num"] for g in group])
    node = copy.deepcopy(template)
    data = {"referencesIds": ids,
            "referencesOptions": {i: dict(OPT) for i in ids},
            "hasBrokenReferences": False, "hasManualEdits": False,
            "isEmpty": False, "citationType": "inline",
            "id": random.randint(-2 ** 31, 2 ** 31 - 1),
            "citationText": SPAN % text}
    pr = node.find(qn("w:sdtPr"))
    pr.find(qn("w:tag")).set(qn("w:val"), json.dumps(data))
    wid = pr.find(qn("w:id"))
    if wid is not None:
        wid.set(qn("w:val"), str(random.randint(10 ** 8, 2 ** 31 - 1)))
    content = node.find(qn("w:sdtContent"))
    runs = content.findall(qn("w:r"))
    for r in runs[1:]:
        content.remove(r)
    t = runs[0].find(qn("w:t"))
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return node, text

TEXT = (
 "Resection extent was reported unevenly across the 27 oncological studies: "
 "rib count was recoverable in 24, defect area in cm² in 10, and a total, "
 "subtotal or partial grading of sternectomy in 8. Study-level mean rib counts "
 "ranged from 2.3 to 4.8, so this literature describes a narrow band of extent "
 "and does not speak to total sternectomy or very large defects. Twelve "
 "studies reported rib count separately by reconstruction arm{C1} and eight "
 "reported defect area at the level of a reconstruction arm.{C2} Across those "
 "study arms, rigid and soft-mesh reconstructions were used for resections of "
 "the same size (median 3.4 versus 3.4 ribs; 146 versus 144 cm²), while "
 "arms receiving no skeletal reconstruction carried smaller resections (median "
 "2.2 ribs; 101 cm², excluding one series whose unreconstructed cases were "
 "soft-tissue resections without bone). The largest oncological series "
 "reported the same pattern within a single cohort — 3.4 ribs resected "
 "with rigid versus 2.7 with flexible reconstruction (p<0.001) — and "
 "described rib count as a surrogate for defect size.{C3}")

d = docx.Document(SRC)
before = d.element.body.xml.count("<w:sdt>")

anchor = [p for p in d.paragraphs
          if "pre-operative antibiotics by one" in annotate(p)][0]
template = citation_nodes(d.paragraphs[159])[2]

new = copy.deepcopy(anchor._p)
for c in list(new.iterchildren()):
    if c.tag in (qn("w:sdt"), qn("w:hyperlink")):
        new.remove(c)
runs = new.findall(qn("w:r"))
for r in runs[1:]:
    new.remove(r)                      # keep one run as the formatting template

for grp in (RIBS, AREA, SPICER):
    node, shown = make_citation(template, grp)
    new.append(node)
    print("  citation:", shown)

anchor._p.addnext(new)
np = Paragraph(new, anchor._parent)
set_text(np, TEXT)

d.save(DST)
chk = docx.Document(DST)
print("sdt before/after:", before, chk.element.body.xml.count("<w:sdt>"))
hit = [p for p in chk.paragraphs if "Resection extent was reported" in p.text][0]
print("\n" + hit.text)
